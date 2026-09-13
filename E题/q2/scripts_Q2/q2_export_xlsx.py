# -*- coding: utf-8 -*-
"""把 Q2 的 CSV 汇总成一个便于人工阅读的 xlsx（多工作表 + 格式美化）。
输出：E题/q2/data_Q2/Q2关键词分类汇总.xlsx
"""
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
Q2 = os.path.dirname(BASE)
D = os.path.join(Q2, "data_Q2")
OUT = os.path.join(D, "Q2关键词分类汇总.xlsx")

def rd(name):
    return pd.read_csv(os.path.join(D, name), encoding="utf-8-sig")

rec = rd("q2_record_clean.csv")
act = rd("q2_active_topsis.csv")
cat = rd("q2_category_summary.csv")
unit = rd("q2_unit_composition.csv")
contrib = rd("q2_contribution.csv")
sens = rd("q2_sensitivity_summary.csv")
rep = rd("q2_representative.csv")
ideal = rd("q2_ideal_solution.csv")
desc = rd("q2_benefit_desc.csv")

CATS = ["黄金词", "重点词", "潜力词", "问题词", "无效词"]
C_THR, S_THR = 8.06, 0.3193
CAT_FILL = {"黄金词": "FFE699", "重点词": "BDD7EE", "潜力词": "C6E0B4",
            "问题词": "F8CBAD", "无效词": "D9D9D9"}

# ---------------- 组装各表 ----------------
# 五类汇总（数量+占比+绝对量+均值）
rows = []
for c in CATS:
    s = rec[rec["基准类别"] == c]
    sa = act[act["基准类别"] == c]
    rows.append({
        "类别": c, "记录数": len(s), "记录数占比%": len(s) / len(rec) * 100,
        "消费额(元)": s["消费额"].sum(), "消费额占比%": s["消费额"].sum() / rec["消费额"].sum() * 100,
        "点击量": s["点击量"].sum(), "点击量占比%": s["点击量"].sum() / rec["点击量"].sum() * 100,
        "浏览量": s["浏览量"].sum(), "浏览量占比%": s["浏览量"].sum() / rec["浏览量"].sum() * 100,
        "单条平均消费(元)": s["消费额"].mean(),
        "单条平均点击": s["点击量"].mean(),
        "平均效益得分S": sa["S"].mean() if len(sa) else np.nan,
    })
sum5 = pd.DataFrame(rows)

# 单元构成：数量 + 占比
uc = (rec.groupby(["方案ID", "推广单元ID", "基准类别"]).size()
        .unstack(fill_value=0).reindex(columns=CATS, fill_value=0).reset_index())
uc.insert(2, "总记录数", uc[CATS].sum(axis=1))
uc_cnt = uc.copy()
uc_pct = uc.copy()
for c in CATS:
    uc_pct[c] = (uc[c] / uc["总记录数"] * 100).round(2)

# 明细：全部记录
det = rec.copy()
det["状态"] = np.where(det["是否活跃"], "活跃", "无效")
det["成本档"] = np.where(det["是否活跃"], np.where(det["消费额"] < C_THR, "低成本", "高成本"), "无成本")
det["效益档"] = np.where(det["是否活跃"], np.where(det["S"] >= S_THR, "高效益", "低效益"), "无效益")
det = det[["序号", "关键词", "方案ID", "推广单元ID", "状态", "成本档", "效益档",
           "消费额", "点击量", "浏览量", "浏览深度", "跳出率", "留存率", "时长秒", "CPC",
           "S", "基准类别", "敏感性"]]
det.columns = ["序号", "关键词", "方案ID", "推广单元ID", "状态", "成本档", "效益档",
               "消费额", "点击量", "浏览量", "浏览深度", "跳出率", "留存率", "时长秒", "CPC",
               "效益得分S", "基准类别", "敏感性"]

det_act = act.copy()
det_act = det_act[["序号", "关键词", "方案ID", "推广单元ID", "消费额", "点击量", "浏览量",
                   "x1", "x2", "x3", "x4", "d_plus", "d_minus", "S",
                   "基准类别", "敏感性"]].copy()
det_act.columns = ["序号", "关键词", "方案ID", "推广单元ID", "消费额", "点击量", "浏览量",
                   "x1 点击规模", "x2 浏览深度", "x3 留存率", "x4 访问时长",
                   "d+ 到正理想解", "d- 到负理想解", "效益得分S", "基准类别", "敏感性"]

# 敏感性
scen_rows = []
for name in ["基准等权", "流量优先", "质量优先"]:
    vc = act["类别_" + name].value_counts().reindex(CATS[:4]).fillna(0).astype(int)
    scen_rows.append({"情景": name, "权重(点击/深度/留存/时长)":
                      {"基准等权": "0.25/0.25/0.25/0.25", "流量优先": "0.40/0.20/0.20/0.20",
                       "质量优先": "0.20/0.25/0.30/0.25"}[name],
                      "S中位数": round(float(np.median(act["S_" + name])), 4),
                      "黄金词": vc["黄金词"], "重点词": vc["重点词"],
                      "潜力词": vc["潜力词"], "问题词": vc["问题词"]})
scen = pd.DataFrame(scen_rows)

# ---------------- 写 Excel ----------------
wb = Workbook()
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10.5)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(color="595959", size=10)

def write_sheet(ws, df, numfmt=None, title=None, subtitle=None, cat_col=None,
                freeze=None, wrap_cols=None, tab_color=None, autofilter=True):
    numfmt = numfmt or {}
    r0 = 1
    if title:
        ws.cell(1, 1, title).font = TITLE_FONT
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(df.columns)))
        ws.row_dimensions[1].height = 22
        r0 = 2
        if subtitle:
            ws.cell(2, 1, subtitle).font = SUB_FONT
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(df.columns)))
            r0 = 3
    hdr = r0
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(hdr, j, c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=hdr + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, (np.floating, float)) and pd.isna(v):
                v = None
            cell = ws.cell(i, j, v)
            cell.border = BORDER
            col = df.columns[j - 1]
            if col in numfmt:
                cell.number_format = numfmt[col]
            if col in (wrap_cols or []):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    # 列宽
    for j, c in enumerate(df.columns, start=1):
        mx = len(str(c))
        for i in range(hdr + 1, hdr + 1 + len(df)):
            val = ws.cell(i, j).value
            if val is not None:
                mx = max(mx, len(f"{val:,.2f}" if isinstance(val, (int, float)) and c in numfmt else str(val)))
        ws.column_dimensions[get_column_letter(j)].width = min(max(mx + 3, 9), 34)
    # 类别着色
    if cat_col and cat_col in df.columns:
        cj = list(df.columns).index(cat_col) + 1
        for i in range(hdr + 1, hdr + 1 + len(df)):
            v = ws.cell(i, cj).value
            if v in CAT_FILL:
                ws.cell(i, cj).fill = PatternFill("solid", fgColor=CAT_FILL[v])
    ws.freeze_panes = freeze or (ws.cell(hdr + 1, 1).coordinate)
    if autofilter and len(df) > 0:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(df.columns))}{hdr+len(df)}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    return hdr

# 0 说明
ws = wb.active; ws.title = "说明"
ws.sheet_properties.tabColor = "1F4E79"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 100
ws["A1"] = "E题 问题2 关键词「成本—效益」五分类 汇总表"; ws["A1"].font = TITLE_FONT
info = [
    ("数据源", "附件1 Sheet3（2227 条“方案ID—推广单元ID—关键词”记录）"),
    ("分类对象", "关键词投放记录（同一关键词跨单元分别分类，不按名称去重）"),
    ("成本指标", "关键词全年消费额 C；阈值 = 活跃记录中位数 8.06 元"),
    ("效益指标", "x1=ln(1+点击量)，x2=ln(1+浏览深度)，x3=1-跳出率，x4=ln(1+时长秒)"),
    ("效益得分", "四指标等权 TOPSIS，S=d-/(d++d-)；阈值 = S 中位数 0.3193"),
    ("五分类", "低成本高效益=黄金词；高成本高效益=重点词；低成本低效益=潜力词；高成本低效益=问题词；无成本无点击=无效词"),
    ("敏感性", "三种权重情景（等权/流量优先/质量优先），三情景类别完全一致为稳定词"),
    ("", ""),
    ("全部记录", f"{len(rec)} 条"),
    ("无效词", f"{(rec['基准类别']=='无效词').sum()} 条（不参与 TOPSIS）"),
    ("活跃记录", f"{(rec['是否活跃']).sum()} 条（进入 TOPSIS）"),
    ("稳定/敏感", f"{int(act['稳定词'].sum())} / {int((~act['稳定词']).sum())} 条，活跃稳定率 {act['稳定词'].mean()*100:.2f}%"),
    ("", ""),
    ("工作表导航", "五类汇总 / 单元构成_数量 / 单元构成_占比 / 各类贡献 / 敏感性检验 / 代表关键词 / 活跃词TOPSIS明细 / 全部记录分类明细 / 理想解与权重 / 效益指标统计"),
    ("result2.xlsx", "按附件模板输出的宽表（另存于同目录）"),
]
r = 3
for k, v in info:
    ws.cell(r, 1, k).font = Font(bold=True, size=10.5)
    ws.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws.freeze_panes = "A3"

# 1 五类汇总（含合计行）
sum5 = pd.concat([sum5, pd.DataFrame([{
    "类别": "合计", "记录数": sum5["记录数"].sum(), "记录数占比%": 100.0,
    "消费额(元)": sum5["消费额(元)"].sum(), "消费额占比%": 100.0,
    "点击量": sum5["点击量"].sum(), "点击量占比%": 100.0,
    "浏览量": sum5["浏览量"].sum(), "浏览量占比%": 100.0,
    "单条平均消费(元)": np.nan, "单条平均点击": np.nan, "平均效益得分S": np.nan}])], ignore_index=True)
write_sheet(wb.create_sheet("五类汇总"), sum5, cat_col="类别", freeze="A2", tab_color="C00000", autofilter=False,
            numfmt={"记录数占比%": "0.00", "消费额(元)": "#,##0.00", "消费额占比%": "0.00",
                    "点击量占比%": "0.00", "浏览量占比%": "0.00",
                    "单条平均消费(元)": "#,##0.00", "单条平均点击": "0.00", "平均效益得分S": "0.0000"})

# 2/3 单元构成（含合计行）
uc_cnt = pd.concat([uc_cnt, pd.DataFrame([{
    "方案ID": "合计", "推广单元ID": "", "总记录数": uc_cnt["总记录数"].sum(),
    **{c: uc_cnt[c].sum() for c in CATS}}])], ignore_index=True)
uc_pct2 = uc_pct.copy()
uc_pct2 = pd.concat([uc_pct2, pd.DataFrame([{
    "方案ID": "合计", "推广单元ID": "", "总记录数": uc_cnt["总记录数"].iloc[:-1].sum(),
    **{c: round(uc_cnt[c].iloc[:-1].sum() / uc_cnt["总记录数"].iloc[:-1].sum() * 100, 2) for c in CATS}}])], ignore_index=True)
write_sheet(wb.create_sheet("单元构成_数量"), uc_cnt, cat_col=None, freeze="A3", tab_color="2E75B6",
            numfmt={c: "#,##0" for c in CATS + ["总记录数"]}, title="各推广单元五类关键词数量", autofilter=False)
write_sheet(wb.create_sheet("单元构成_占比"), uc_pct2, freeze="A3", tab_color="2E75B6",
            numfmt={c: "0.00" for c in CATS}, title="各推广单元五类关键词占比（%）", autofilter=False)

# 4 各类贡献
write_sheet(wb.create_sheet("各类贡献"), contrib, cat_col="类别", freeze="A2", tab_color="548235", autofilter=False,
            numfmt={"消费额占比": "0.00", "点击量占比": "0.00", "浏览量占比": "0.00"})

# 5 敏感性
write_sheet(wb.create_sheet("敏感性检验"), scen, freeze="A3", tab_color="7030A0",
            numfmt={"S中位数": "0.0000"}, title="三种权重情景下的分类稳定性")
ws = wb["敏感性检验"]
start = 2 + 1 + len(scen) + 1
ws.cell(start, 1, "稳定性汇总").font = Font(bold=True, size=11, color="1F4E79")
for i, row in enumerate(sens.itertuples(index=False), start=start + 1):
    ws.cell(i, 1, row[0]).font = Font(bold=True)
    ws.cell(i, 2, row[1])

# 6 代表关键词
write_sheet(wb.create_sheet("代表关键词"), rep, cat_col="类别", freeze="A2", tab_color="BF8F00",
            numfmt={"消费额": "#,##0.00", "浏览深度": "0.0000", "留存率": "0.0000", "时长秒": "0", "S": "0.0000"})

# 7 活跃词 TOPSIS 明细
write_sheet(wb.create_sheet("活跃词TOPSIS明细"), det_act, cat_col="基准类别", freeze="A2", tab_color="00B050",
            numfmt={"消费额": "#,##0.00", "x1 点击规模": "0.0000", "x2 浏览深度": "0.0000",
                    "x3 留存率": "0.0000", "x4 访问时长": "0.0000",
                    "d+ 到正理想解": "0.0000", "d- 到负理想解": "0.0000", "效益得分S": "0.0000"})

# 8 全部记录明细
write_sheet(wb.create_sheet("全部记录分类明细"), det, cat_col="基准类别", freeze="A2", tab_color="808080",
            numfmt={"消费额": "#,##0.00", "浏览深度": "0.0000", "跳出率": "0.0000", "留存率": "0.0000",
                    "时长秒": "0", "CPC": "0.0000", "效益得分S": "0.0000"})

# 9 理想解与权重
write_sheet(wb.create_sheet("理想解与权重"), ideal, freeze="A3", tab_color="ED7D31",
            numfmt={"权重": "0.00", "正理想解v+": "0.000000", "负理想解v-": "0.000000"},
            title="TOPSIS 正/负理想解（加权标准化值）")

# 10 效益指标统计
write_sheet(wb.create_sheet("效益指标统计"), desc, freeze="A3", tab_color="ED7D31",
            numfmt={c: "0.0000" for c in ["min", "25%", "50%", "75%", "max", "mean"]},
            title="效益指标描述统计（1337 条活跃记录，缩尾+对数变换后）")

wb.save(OUT)
print("已保存:", OUT)
print("工作表:", wb.sheetnames)
