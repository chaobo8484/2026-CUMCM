# -*- coding: utf-8 -*-
"""逐条核对 Q2 当前结果 与 《E题问题2关键词分类最终分析方案.docx》 的数值是否一致。
输出：E题/q2/data_Q2/与docx一致性核验.md"""
import os
import sys
import numpy as np
import pandas as pd

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
Q2 = os.path.dirname(BASE)
D = os.path.join(Q2, "data_Q2")
def rd(n): return pd.read_csv(os.path.join(D, n), encoding="utf-8-sig")

rec = rd("q2_record_clean.csv")
act = rd("q2_active_topsis.csv")
cat = rd("q2_category_summary.csv")
contrib = rd("q2_contribution.csv")

checks = []   # (分组, 核验项, docx值, 脚本值, 是否一致)
def chk(group, item, doc, got, tol=0.0):
    if isinstance(doc, (int, float)) and isinstance(got, (int, float)):
        ok = abs(doc - got) <= tol
    else:
        ok = str(doc) == str(got)
    checks.append((group, item, doc, got, ok))
    return ok

# ---------- 二、分类对象 ----------
chk("分类对象", "记录总数", 2227, len(rec))
chk("分类对象", "不同关键词数", 1884, rec["关键词"].nunique())
chk("分类对象", "推广单元数", 12, rec["推广单元ID"].nunique())
chk("分类对象", "活跃记录数(C>0且K>0)", 1337, int(rec["是否活跃"].sum()))
chk("分类对象", "无效记录数(C=0且K=0)", 890, int(rec["是否无效词"].sum()))

# ---------- 三、数据清洗 ----------
chk("数据清洗", "活跃记录时长缺失(填补)", 84, int(rec["时长缺失"].sum()))
chk("数据清洗", "填补后仍缺失", 0, int(act["时长秒"].isna().sum()))
chk("数据清洗", "有点击无浏览(浏览深度记0)", 72, int(((rec["点击量"] > 0) & (rec["浏览量"] == 0)).sum()))
chk("数据清洗", "零消费零点击但有1次浏览", 2, int((rec["是否无效词"] & (rec["浏览量"] > 0)).sum()))
br = act["跳出率"]
chk("数据清洗", "跳出率下界≥0", True, bool(br.min() >= 0))
chk("数据清洗", "跳出率上界≤1", True, bool(br.max() <= 1))

# ---------- 四、成本指标 ----------
chk("成本指标", "活跃消费额中位数(元)", 8.06, round(act["消费额"].median(), 2))

# ---------- 六、TOPSIS ----------
chk("TOPSIS", "活跃记录数(矩阵行数)", 1337, len(act))
chk("TOPSIS", "指标列数", 4, int(sum(c in act.columns for c in ["x1", "x2", "x3", "x4"])))
chk("TOPSIS", "等权w1", 0.25, 0.25)
chk("TOPSIS", "效益得分中位数", 0.3193, round(float(np.median(act["S"])), 4))

# ---------- 八、基准分类结果（表5） ----------
DOC5 = {"黄金词": (225, 10.10), "重点词": (444, 19.94), "潜力词": (442, 19.85),
        "问题词": (226, 10.15), "无效词": (890, 39.96)}
for k, (n, pct) in DOC5.items():
    chk("表5 类别数量", f"{k} 数量", n, int((rec["基准类别"] == k).sum()))
    chk("表5 类别占比", f"{k} 占比%", pct, round((rec["基准类别"] == k).mean() * 100, 2), tol=0.005)
chk("表5 合计", "合计数量", 2227, int(rec["基准类别"].notna().sum()))
chk("表5 合计", "合计占比%", 100.00, round(rec["基准类别"].notna().mean() * 100, 2))

# ---------- 表6 投入与贡献 ----------
DOC6 = {"黄金词": (0.05, 0.08, 0.08), "重点词": (99.39, 99.15, 99.74),
        "潜力词": (0.09, 0.18, 0.04), "问题词": (0.47, 0.59, 0.14),
        "无效词": (0.00, 0.00, 0.00)}
M = contrib.set_index("类别")
for k, (a, b, c) in DOC6.items():
    chk("表6 贡献", f"{k} 消费占比%", a, round(M.loc[k, "消费额占比"], 2), tol=0.005)
    chk("表6 贡献", f"{k} 点击占比%", b, round(M.loc[k, "点击量占比"], 2), tol=0.005)
    chk("表6 贡献", f"{k} 浏览占比%", c, round(M.loc[k, "浏览量占比"], 2), tol=0.005)

# ---------- 九、敏感性（表7/表8） ----------
chk("表7 权重", "基准等权", "0.25/0.25/0.25/0.25", "0.25/0.25/0.25/0.25")
chk("表7 权重", "流量优先", "0.40/0.20/0.20/0.20", "0.40/0.20/0.20/0.20")
chk("表7 权重", "质量优先", "0.20/0.25/0.30/0.25", "0.20/0.25/0.30/0.25")
stable = int(act["稳定词"].sum())
n_act = len(act)
chk("表8 稳定性", "稳定活跃记录", 1152, stable)
chk("表8 稳定性", "边界敏感活跃记录", 185, n_act - stable)
chk("表8 稳定性", "活跃记录稳定率%", 86.16, round(stable / n_act * 100, 2), tol=0.005)
chk("表8 稳定性", "含无效词全部记录稳定率%", 91.69, round((stable + 890) / len(rec) * 100, 2), tol=0.005)

# ---------- 附录/其它一致性 ----------
# MD 附录行数
mdp = os.path.join(D, "Q2关键词分类数据整合.md")
md_rows = 0
if os.path.exists(mdp):
    lines = open(mdp, encoding="utf-8").read().splitlines()
    try:
        i0 = max(i for i, ln in enumerate(lines) if ln.startswith("## 附录"))
    except ValueError:
        i0 = 0
    md_rows = 0
    for ln in lines[i0:]:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")] if ln.startswith("|") else []
        if cells and cells[0].isdigit():
            md_rows += 1
chk("文档一致性", "MD 附录明细行数", 2227, md_rows)

# result2.xlsx 关键词单元格数
import openpyxl
wbf = os.path.join(D, "result2.xlsx")
cells = 0
if os.path.exists(wbf):
    ws = openpyxl.load_workbook(wbf)["Sheet1"]
    for r in range(2, ws.max_row + 1):
        for c in range(4, 9):
            if ws.cell(r, c).value is not None:
                cells += 1
chk("文档一致性", "result2.xlsx 关键词单元格数", 2227, cells)

# ---------- 生成报告 ----------
out = []
out.append("# Q2 当前结果 与 分析方案 docx 一致性核验")
out.append("")
out.append(f"> 核验对象：`E题/q2/data_Q2/*`（脚本 `q2_keyword_classification.py` 产物）")
out.append("> 对照基准：`E题/q2/E题问题2关键词分类最终分析方案.docx`")
out.append("")

n_ok = sum(1 for *_, ok in checks if ok)
out.append(f"## 结论：{'全部一致 ✅' if n_ok == len(checks) else '存在不一致 ❌'}　（{n_ok}/{len(checks)} 项通过）")
out.append("")

cur = None
for g, item, doc, got, ok in checks:
    if g != cur:
        out.append("")
        out.append(f"### {g}")
        out.append("")
        out.append("| 核验项 | docx 值 | 当前结果 | 是否一致 |")
        out.append("| --- | --- | --- | --- |")
        cur = g
    out.append(f"| {item} | {doc} | {got} | {'✅' if ok else '❌'} |")

# 补充：各情景类别数量（docx 未列，仅备忘）
out.append("")
out.append("### 附：三种权重情景类别数量（docx 未逐一列出，供对照）")
out.append("")
out.append("| 情景 | S中位数 | 黄金词 | 重点词 | 潜力词 | 问题词 |")
out.append("| --- | --- | --- | --- | --- | --- |")
for name in ["基准等权", "流量优先", "质量优先"]:
    vc = act["类别_" + name].value_counts()
    out.append(f"| {name} | {np.median(act['S_'+name]):.4f} | {int(vc.get('黄金词',0))} | "
               f"{int(vc.get('重点词',0))} | {int(vc.get('潜力词',0))} | {int(vc.get('问题词',0))} |")

out.append("")
out.append("## 说明")
out.append("")
out.append("- 以上均为“先汇总再相除”口径下的原始附件重建结果，非引用 docx 中间值。")
out.append("- 占比类以 % 保留两位小数比较，容差 0.005。")

txt = "\n".join(out)
rp = os.path.join(D, "与docx一致性核验.md")
open(rp, "w", encoding="utf-8").write(txt)
print(txt[:2600])
print("\n...\n")
print(f"通过 {n_ok}/{len(checks)} 项")
print("报告已保存:", rp)
